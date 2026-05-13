import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from django.core.paginator import Paginator
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.db.models import Count, Q
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from datetime import datetime, date, timedelta
import calendar
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, cm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, PageBreak, Image)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.pdfbase.pdfmetrics import stringWidth
from .models import Ticket, Configuracion, Notificacion, NotificacionUsuario, AuditLog, Tipificacion
from .forms import TicketForm, TicketUpdateForm, StaffTicketForm
from .ai_analyzer import analizar as analizar_ticket


def _crear_notificacion(mensaje, url='', usuario=None):
    Notificacion.objects.create(mensaje=mensaje, url=url, usuario=usuario)


def _notificar_admins(mensaje, url=''):
    from django.contrib.auth.models import User
    admins = User.objects.filter(groups__name='Administrador')
    for admin in admins:
        Notificacion.objects.create(mensaje=mensaje, url=url, usuario=admin)


def _crear_notificacion_usuario(mensaje, url='', solicitante='', direccion_area=''):
    NotificacionUsuario.objects.create(
        mensaje=mensaje, url=url,
        solicitante=solicitante, direccion_area=direccion_area
    )


def _get_client_ip(request):
    real_ip = request.COOKIES.get('real_ip')
    if real_ip:
        return real_ip
    x_forwarded = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded:
        ip = x_forwarded.split(',')[0].strip()
        if ip:
            return ip
    x_real_ip = request.META.get('HTTP_X_REAL_IP')
    if x_real_ip:
        return x_real_ip
    return request.META.get('REMOTE_ADDR', '0.0.0.0')


def _registrar_auditoria(request, usuario, accion, modelo, id_objeto=None, detalles='', nombre_completo=None, direccion_area=''):
    ip = _get_client_ip(request)
    if not nombre_completo:
        nombre_completo = usuario.get_full_name() if usuario else 'Anónimo'
    if not direccion_area:
        if usuario and usuario.is_authenticated:
            if usuario.groups.filter(name__in=['Administrador', 'Tecnico']).exists():
                direccion_area = 'Tecnología'
        else:
            direccion_area = request.session.get('usuario_area', '')
    AuditLog.objects.create(
        usuario=usuario,
        nombre_completo=nombre_completo,
        accion=accion,
        modelo=modelo,
        id_objeto=id_objeto,
        detalles=detalles,
        direccion_ip=ip,
        direccion_area=direccion_area,
    )


def asignar_tecnico_round_robin():
    from django.contrib.auth.models import User
    tecnicos = list(User.objects.filter(groups__name='Tecnico').order_by('id'))
    if not tecnicos:
        return None
    config, _ = Configuracion.objects.get_or_create(
        clave='round_robin_index', defaults={'valor': 0}
    )
    idx = config.valor % len(tecnicos)
    config.valor = idx + 1
    config.save(update_fields=['valor'])
    return tecnicos[idx]


def is_admin(user):
    return user.is_authenticated and user.groups.filter(name='Administrador').exists()


def is_tecnico(user):
    return user.is_authenticated and user.groups.filter(name='Tecnico').exists()


def is_admin_or_tecnico(user):
    return is_admin(user) or is_tecnico(user)


def landing(request):
    return render(request, 'tickets/landing.html')


def _redirect_if_staff(request):
    if request.user.is_authenticated and is_admin_or_tecnico(request.user):
        return True
    return False


def crear_ticket(request):
    if _redirect_if_staff(request):
        return redirect('dashboard')
    nombre = request.session.get('usuario_nombre')
    area = request.session.get('usuario_area')

    if not nombre or not area:
        return redirect('landing')

    if request.method == 'POST':
        form = TicketForm(request.POST)
        if form.is_valid():
            ticket = form.save(commit=False)
            analisis = analizar_ticket(ticket.descripcion)
            ticket.save()
            if analisis and analisis.get('tipos_detectados'):
                for slug in analisis['tipos_detectados']:
                    tip, _ = Tipificacion.objects.get_or_create(
                        slug=slug,
                        defaults={'nombre': dict(Ticket.TIPO_FALLA_CHOICES).get(slug, slug)}
                    )
                    ticket.tipificaciones.add(tip)
            tiene_tipo_automatico = ticket.tipificaciones.exclude(slug='otro').exists()
            if not ticket.tecnico_asignado and tiene_tipo_automatico:
                tecnico = asignar_tecnico_round_robin()
                if tecnico:
                    ticket.tecnico_asignado = tecnico
                    ticket.save(update_fields=['tecnico_asignado'])
            if ticket.tecnico_asignado:
                _crear_notificacion(
                    f'Nuevo ticket {ticket.codigo} asignado a usted - {ticket.solicitante}',
                    url=f'/tickets/{ticket.pk}/',
                    usuario=ticket.tecnico_asignado
                )
                _notificar_admins(
                    f'Nuevo ticket {ticket.codigo} - {ticket.solicitante} asignado a {ticket.nombre_tecnico}',
                    url=f'/tickets/{ticket.pk}/'
                )
            else:
                _notificar_admins(
                    f'Nuevo ticket {ticket.codigo} - {ticket.solicitante} (Solicitud de servicio)',
                    url=f'/tickets/{ticket.pk}/'
                )
            _registrar_auditoria(request, None, 'CREAR', 'Ticket', id_objeto=ticket.pk, detalles=f'Usuario "{nombre}" creó el ticket {ticket.codigo} - {ticket.descripcion[:100]}', nombre_completo=nombre)
            messages.success(request, f'Ticket creado exitosamente. Codigo: {ticket.codigo}')
            return redirect('panel_usuario')
    else:
        form = TicketForm(initial={
            'solicitante': nombre,
            'direccion_area': area,
        })
    return render(request, 'tickets/crear_ticket.html', {'form': form})


def panel_usuario_login(request):
    if _redirect_if_staff(request):
        return redirect('dashboard')
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        area = request.POST.get('area', '').strip()
        if nombre and area:
            request.session['usuario_nombre'] = nombre
            request.session['usuario_area'] = area
            _registrar_auditoria(request, None, 'ACCESO', 'Sesión', detalles=f'Acceso al panel de usuario - {nombre}', nombre_completo=nombre, direccion_area=area)
            return redirect('panel_usuario')
        else:
            messages.error(request, 'Debe ingresar nombre y area.')
    return render(request, 'tickets/panel_usuario_login.html')


def panel_usuario(request):
    if _redirect_if_staff(request):
        return redirect('dashboard')
    nombre = request.session.get('usuario_nombre')
    area = request.session.get('usuario_area')

    if not nombre or not area:
        return redirect('landing')

    mis_tickets = Ticket.objects.filter(
        solicitante__iexact=nombre,
        direccion_area__iexact=area
    ).order_by('-fecha_creacion')

    paginator = Paginator(mis_tickets, 10)
    page = request.GET.get('page')
    mis_tickets_page = paginator.get_page(page)

    context = {
        'nombre': nombre,
        'area': area,
        'mis_tickets': mis_tickets_page,
        'total': mis_tickets.count(),
        'abiertos': mis_tickets.filter(estado='abierto').count(),
        'en_proceso': mis_tickets.filter(estado='en_proceso').count(),
        'pendientes': mis_tickets.filter(estado='pendiente').count(),
        'resueltos': mis_tickets.filter(estado='resuelto').count(),
    }
    return render(request, 'tickets/panel_usuario.html', context)


def cerrar_sesion_usuario(request):
    nombre = request.session.get('usuario_nombre', '')
    area = request.session.get('usuario_area', '')
    request.session.pop('usuario_nombre', None)
    request.session.pop('usuario_area', None)
    _registrar_auditoria(request, None, 'LOGOUT', 'Sesión', detalles=f'Cierre de sesión de usuario - {nombre}', nombre_completo=nombre, direccion_area=area)
    return redirect('landing')


def ver_ticket_usuario(request, ticket_id):
    if _redirect_if_staff(request):
        return redirect('dashboard')
    nombre = request.session.get('usuario_nombre')
    area = request.session.get('usuario_area')

    if not nombre or not area:
        return redirect('landing')

    ticket = get_object_or_404(Ticket, pk=ticket_id)

    if ticket.solicitante.strip().lower() != nombre.lower() or ticket.direccion_area.strip().lower() != area.lower():
        messages.error(request, 'No tiene permiso para ver este ticket.')
        return redirect('panel_usuario')

    return render(request, 'tickets/ver_ticket_usuario.html', {'ticket': ticket})


@login_required
@user_passes_test(is_admin_or_tecnico)
def dashboard(request):
    tickets = Ticket.objects.all().select_related('tecnico_asignado')

    if is_tecnico(request.user):
        tickets = tickets.filter(tecnico_asignado=request.user)

    filtro_estado = request.GET.get('estado')
    filtro_tecnico = request.GET.get('tecnico')
    filtro_tipo_falla = request.GET.get('tipo_falla')
    buscar = request.GET.get('buscar')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    if filtro_estado == '_sin_asignar':
        tickets = tickets.filter(tecnico_asignado__isnull=True)
    elif filtro_estado:
        tickets = tickets.filter(estado=filtro_estado)
    if filtro_tecnico:
        tickets = tickets.filter(tecnico_asignado_id=filtro_tecnico)
    if filtro_tipo_falla:
        tickets = tickets.filter(tipificaciones__slug=filtro_tipo_falla)
    if buscar:
        tickets = tickets.filter(
            Q(solicitante__icontains=buscar) |
            Q(descripcion__icontains=buscar) |
            Q(direccion_area__icontains=buscar)
        )
    if fecha_desde:
        tickets = tickets.filter(fecha_creacion__gte=fecha_desde)
    if fecha_hasta:
        tickets = tickets.filter(fecha_creacion__lte=fecha_hasta)

    estadisticas = {
        'total': tickets.count(),
        'abiertos': tickets.filter(estado='abierto').count(),
        'en_proceso': tickets.filter(estado='en_proceso').count(),
        'pendientes': tickets.filter(estado='pendiente').count(),
        'resueltos': tickets.filter(estado='resuelto').count(),
    }

    tecnicos = None
    if is_admin(request.user):
        from django.contrib.auth.models import User
        tecnicos = User.objects.filter(groups__name='Tecnico')

    paginator = Paginator(tickets, 10)
    page = request.GET.get('page')
    tickets_page = paginator.get_page(page)

    context = {
        'tickets': tickets_page,
        'estadisticas': estadisticas,
        'tecnicos': tecnicos,
        'es_tecnico': is_tecnico(request.user),
        'es_admin': is_admin(request.user),
    }
    return render(request, 'tickets/dashboard.html', context)


@login_required
@user_passes_test(is_admin_or_tecnico)
def detalle_ticket(request, ticket_id):
    ticket = get_object_or_404(Ticket, pk=ticket_id)
    es_admin = is_admin(request.user)
    tecnico_original_id = ticket.tecnico_asignado_id
    estado_original = ticket.estado
    notas_original = ticket.notas
    tipificaciones_original = set(ticket.tipificaciones.values_list('slug', flat=True))

    if request.method == 'POST':
        form = TicketUpdateForm(request.POST, instance=ticket)
        if form.is_valid():
            if not es_admin:
                form.instance.tecnico_asignado_id = tecnico_original_id
            elif request.POST.get('atendere_yo'):
                form.instance.tecnico_asignado = request.user
            nuevo_estado = form.cleaned_data.get('estado')
            if estado_original != nuevo_estado:
                messages.info(request, f'Estado cambiado a: {dict(Ticket.ESTADO_CHOICES).get(nuevo_estado)}')
            form.save()
            nombre_user = request.user.get_full_name() or request.user.username
            if estado_original != nuevo_estado:
                if es_admin and ticket.tecnico_asignado:
                    _crear_notificacion(
                        f'{ticket.codigo} cambió a {dict(Ticket.ESTADO_CHOICES).get(nuevo_estado)} por {nombre_user}',
                        url=f'/tickets/{ticket.pk}/',
                        usuario=ticket.tecnico_asignado
                    )
                elif not es_admin:
                    _notificar_admins(
                        f'{ticket.codigo} cambió a {dict(Ticket.ESTADO_CHOICES).get(nuevo_estado)} por {nombre_user}',
                        url=f'/tickets/{ticket.pk}/'
                    )
                _crear_notificacion_usuario(
                    f'{ticket.codigo} cambió a {dict(Ticket.ESTADO_CHOICES).get(nuevo_estado)}',
                    url=f'/mi-panel/ticket/{ticket.pk}/',
                    solicitante=ticket.solicitante,
                    direccion_area=ticket.direccion_area
                )
            if tecnico_original_id != ticket.tecnico_asignado_id and ticket.tecnico_asignado:
                _crear_notificacion(
                    f'{ticket.codigo} asignado a {ticket.nombre_tecnico} por {nombre_user}',
                    url=f'/tickets/{ticket.pk}/',
                    usuario=ticket.tecnico_asignado
                )
            nuevas_tipificaciones = set(ticket.tipificaciones.values_list('slug', flat=True))
            detalles_cambios = []
            if estado_original != nuevo_estado:
                detalles_cambios.append(f'Estado: {dict(Ticket.ESTADO_CHOICES).get(estado_original)} → {dict(Ticket.ESTADO_CHOICES).get(nuevo_estado)}')
            if tecnico_original_id != ticket.tecnico_asignado_id:
                if ticket.tecnico_asignado:
                    detalles_cambios.append(f'Técnico asignado: {ticket.nombre_tecnico}')
                else:
                    detalles_cambios.append('Técnico removido')
            if notas_original != ticket.notas:
                detalles_cambios.append('Notas de seguimiento editadas')
            if tipificaciones_original != nuevas_tipificaciones:
                agregadas = nuevas_tipificaciones - tipificaciones_original
                quitadas = tipificaciones_original - nuevas_tipificaciones
                partes = []
                if agregadas:
                    nombres_agregadas = list(Tipificacion.objects.filter(slug__in=agregadas).values_list('nombre', flat=True))
                    partes.append(f'agregadas: {", ".join(nombres_agregadas)}')
                if quitadas:
                    nombres_quitadas = list(Tipificacion.objects.filter(slug__in=quitadas).values_list('nombre', flat=True))
                    partes.append(f'quitadas: {", ".join(nombres_quitadas)}')
                detalles_cambios.append(f'Tipificación modificada ({"; ".join(partes)})')
            _registrar_auditoria(request, request.user, 'EDITAR', 'Ticket', id_objeto=ticket.pk, detalles=' | '.join(detalles_cambios) if detalles_cambios else 'Actualización del ticket')
            messages.success(request, 'Ticket actualizado correctamente.')
            return redirect('detalle_ticket', ticket_id=ticket.pk)
    else:
        form = TicketUpdateForm(instance=ticket)

    analisis = analizar_ticket(ticket.descripcion)
    if analisis and analisis.get('tipos_detectados'):
        solo_otro = not ticket.tipificaciones.exclude(slug='otro').exists()
        if not ticket.tipificaciones.exists() or solo_otro:
            for slug in analisis['tipos_detectados']:
                tip, _ = Tipificacion.objects.get_or_create(
                    slug=slug,
                    defaults={'nombre': dict(Ticket.TIPO_FALLA_CHOICES).get(slug, slug)}
                )
                if not ticket.tipificaciones.filter(slug=slug).exists():
                    ticket.tipificaciones.add(tip)

    context = {
        'ticket': ticket,
        'form': form,
        'es_admin': es_admin,
        'analisis': analisis,
    }
    return render(request, 'tickets/detalle_ticket.html', context)


@login_required
@user_passes_test(is_admin_or_tecnico)
def reportes(request):
    from urllib.parse import urlencode
    from .chart_utils import svg_donut, svg_bar_vertical, svg_bar_horizontal, svg_bar_3d, svg_lollipop, svg_lollipop_vertical, ESTADO_LABELS, ESTADO_COLORS, TIPO_COLORS
    import calendar

    # --- Leer filtros activos ---
    filtro_estado = request.GET.get('estado', '')
    filtro_tipo = request.GET.get('tipo', '')
    filtro_direccion = request.GET.get('direccion', '')
    filtro_tecnico_id = request.GET.get('tecnico_id', '')
    filtro_mes = request.GET.get('mes', '')
    filtro_anio = request.GET.get('anio', '')
    periodo_activo = request.GET.get('periodo', 'trimestre')
    if periodo_activo not in ('trimestre', 'semestre', 'año'):
        periodo_activo = 'trimestre'

    # --- Construir base de tickets ---
    tickets = Ticket.objects.all()
    if is_tecnico(request.user):
        tickets = tickets.filter(tecnico_asignado=request.user)

    if filtro_estado:
        tickets = tickets.filter(estado=filtro_estado)
    if filtro_tipo:
        tickets = tickets.filter(tipificaciones__slug=filtro_tipo)
    if filtro_direccion:
        tickets = tickets.filter(direccion_area=filtro_direccion)
    if filtro_tecnico_id:
        tickets = tickets.filter(tecnico_asignado_id=filtro_tecnico_id)
    if filtro_mes and filtro_anio:
        tickets = tickets.filter(
            fecha_creacion__year=int(filtro_anio),
            fecha_creacion__month=int(filtro_mes)
        )
    elif filtro_anio:
        tickets = tickets.filter(fecha_creacion__year=int(filtro_anio))

    # --- Ventana actual de filtros (para construir links) ---
    filtros_actuales = {}
    if filtro_estado: filtros_actuales['estado'] = filtro_estado
    if filtro_tipo: filtros_actuales['tipo'] = filtro_tipo
    if filtro_direccion: filtros_actuales['direccion'] = filtro_direccion
    if filtro_tecnico_id: filtros_actuales['tecnico_id'] = filtro_tecnico_id
    if filtro_mes: filtros_actuales['mes'] = filtro_mes
    if filtro_anio: filtros_actuales['anio'] = filtro_anio
    if periodo_activo: filtros_actuales['periodo'] = periodo_activo

    def filter_url(**kw):
        p = dict(filtros_actuales)
        for k in ('periodo',):  # periodo se maneja aparte
            pass
        for k in list(kw.keys()):
            p.pop(k, None)
        for k, v in kw.items():
            if v:
                p[k] = v
            else:
                p.pop(k, None)
        if not p:
            return '.'
        return '?' + urlencode(p)

    def toggle_url(key, value):
        if filtros_actuales.get(key) == value:
            return filter_url(**{key: ''})
        return filter_url(**{key: value})

    meses_es = ['', 'Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

    # --- Comparación de rangos de fechas ---
    hoy_dt = timezone.now()
    r1_desde_str = request.GET.get('r1_desde', '')
    r1_hasta_str = request.GET.get('r1_hasta', '')
    r2_desde_str = request.GET.get('r2_desde', '')
    r2_hasta_str = request.GET.get('r2_hasta', '')

    def parse_range(desde_str, hasta_str, default_mes_offset, default_meses):
        if desde_str and hasta_str:
            try:
                d = datetime.strptime(desde_str, '%Y-%m-%d').date()
                h = datetime.strptime(hasta_str, '%Y-%m-%d').date()
                return d, h
            except ValueError:
                pass
        m = hoy_dt.month - default_mes_offset
        y = hoy_dt.year
        while m < 1:
            m += 12
            y -= 1
        inicio = date(y, m, 1)
        m2 = m + default_meses - 1
        y2 = y
        while m2 > 12:
            m2 -= 12
            y2 += 1
        ultimo = calendar.monthrange(y2, m2)[1]
        fin = date(y2, m2, ultimo)
        return inicio, fin

    # --- Detectar rangos activos ---
    r1_explicito = bool(r1_desde_str and r1_hasta_str)
    r2_explicito = bool(r2_desde_str and r2_hasta_str)

    if r1_explicito and r2_explicito:
        modo_rangos = 'compare'
    elif r1_explicito or r2_explicito:
        modo_rangos = 'single'
    else:
        modo_rangos = 'none'

    # Solo computar rangos cuando hay parámetros explícitos
    if r1_explicito:
        r1_desde, r1_hasta = parse_range(r1_desde_str, r1_hasta_str, 3, 3)
    else:
        r1_desde = r1_hasta = None

    if r2_explicito:
        r2_desde, r2_hasta = parse_range(r2_desde_str, r2_hasta_str, 6, 3)
    else:
        r2_desde = r2_hasta = None

    # Rango principal para gráficas
    if modo_rangos == 'none':
        pri_desde = pri_hasta = None
    elif modo_rangos == 'single' and r2_explicito and not r1_explicito:
        pri_desde, pri_hasta = r2_desde, r2_hasta
    else:
        pri_desde, pri_hasta = r1_desde, r1_hasta

    # --- Aplicar rango principal como filtro para gráficas ---
    tickets_base = tickets
    if pri_desde and pri_hasta:
        tickets = tickets.filter(fecha_creacion__gte=pri_desde, fecha_creacion__lte=pri_hasta)

    # --- Tickets por Estado ---
    por_estado = list(tickets.values('estado').annotate(cantidad=Count('id')).order_by('-cantidad'))
    estado_data = [d['cantidad'] for d in por_estado]
    estado_lbl = [ESTADO_LABELS.get(d['estado'], d['estado']) for d in por_estado]
    estado_col = [ESTADO_COLORS.get(d['estado'], '#64748b') for d in por_estado]
    estado_hrefs = [toggle_url('estado', d['estado']) for d in por_estado]
    svg_estado = svg_donut(estado_data, estado_lbl, estado_col, w=480, h=400, inner_r=0, explode=15, depth=12, hrefs=estado_hrefs)

    # --- Tickets por Tipo de Falla ---
    por_tipo_qs = Tipificacion.objects.filter(ticket__in=tickets).annotate(cantidad=Count('ticket')).order_by('-cantidad')
    por_tipo = [{'tipo_falla': t.nombre, 'slug': t.slug, 'cantidad': t.cantidad} for t in por_tipo_qs]
    tipo_data = [d['cantidad'] for d in por_tipo]
    tipo_lbl = [d['tipo_falla'] for d in por_tipo]
    tipo_col = [TIPO_COLORS.get(d['tipo_falla'], '#64748b') for d in por_tipo]
    tipo_hrefs = [toggle_url('tipo', d['slug']) for d in por_tipo]
    svg_tipo = svg_bar_3d(tipo_data, tipo_lbl, tipo_col, hrefs=tipo_hrefs)

    # --- Tickets por Dirección ---
    por_area = list(tickets.values('direccion_area').annotate(cantidad=Count('id')).order_by('direccion_area'))
    area_data = [d['cantidad'] for d in por_area]
    area_lbl = [d['direccion_area'] or '(Sin dirección)' for d in por_area]
    area_hrefs = [toggle_url('direccion', d['direccion_area'] or '') for d in por_area]
    svg_area = svg_lollipop(area_data, area_lbl, '#3347A0', w=780, h=420, pad_l=220, pad_r=60, hrefs=area_hrefs)

    # --- Tickets por Técnico ---
    por_tecnico = list(tickets.filter(
        tecnico_asignado__isnull=False
    ).exclude(
        tecnico_asignado__groups__name='Administrador'
    ).values(
        'tecnico_asignado_id', 'tecnico_asignado__username', 'tecnico_asignado__first_name', 'tecnico_asignado__last_name'
    ).annotate(cantidad=Count('id')).order_by('-cantidad')[:10])
    for item in por_tecnico:
        nombre = f"{item['tecnico_asignado__first_name']} {item['tecnico_asignado__last_name']}".strip()
        item['tecnico_asignado__nombre'] = nombre if nombre else item['tecnico_asignado__username']
    tec_data = [d['cantidad'] for d in por_tecnico]
    tec_lbl = [d['tecnico_asignado__nombre'] for d in por_tecnico]
    tec_hrefs = [toggle_url('tecnico_id', str(d['tecnico_asignado_id'])) for d in por_tecnico]
    svg_tecnico = svg_lollipop_vertical(tec_data, tec_lbl, '#5A63F7', w=600, h=350, pad_l=40, pad_r=20, pad_t=40, pad_b=80, hrefs=tec_hrefs)

    def nav_url_explicit(d, h, r, direccion):
        meses_diff = (h.year - d.year) * 12 + (h.month - d.month) + 1
        nd = d.replace(day=1)
        nm = nd.month + direccion
        ny = nd.year
        while nm < 1:
            nm += 12
            ny -= 1
        while nm > 12:
            nm -= 12
            ny += 1
        nd2 = date(ny, nm, 1)
        nm2 = nm + meses_diff - 1
        ny2 = ny
        while nm2 > 12:
            nm2 -= 12
            ny2 += 1
        while nm2 < 1:
            nm2 += 12
            ny2 -= 1
        ultimo = calendar.monthrange(ny2, nm2)[1]
        nh = date(ny2, nm2, ultimo)
        key = f'r{r}_desde'
        val = nd2.strftime('%Y-%m-%d')
        p = dict(filtros_actuales)
        p[key] = val
        p[f'r{r}_hasta'] = nh.strftime('%Y-%m-%d')
        for k in list(p.keys()):
            if k.startswith('r') and k[1].isdigit() and k not in (f'r{r}_desde', f'r{r}_hasta'):
                p.pop(k, None)
        return '?' + urlencode(p) if p else '.'

    def rango_stats(desde, hasta):
        if desde is None or hasta is None:
            return None
        qs = tickets_base.filter(fecha_creacion__gte=desde, fecha_creacion__lte=hasta)
        total = qs.count()
        por_estado = list(qs.values('estado').annotate(cantidad=Count('id')).order_by('-cantidad'))
        for item in por_estado:
            item['color'] = ESTADO_COLORS.get(item['estado'], '#64748b')
            item['label'] = ESTADO_LABELS.get(item['estado'], item['estado'])
        return {
            'total': total,
            'desde': desde,
            'hasta': hasta,
            'etiqueta': f'{desde.strftime("%d/%m/%Y")} - {hasta.strftime("%d/%m/%Y")}',
            'por_estado': por_estado,
        }

    def rango_seguro(data):
        return data or {'total': 0, 'desde': None, 'hasta': None, 'etiqueta': '', 'por_estado': []}

    r1_data = rango_seguro(rango_stats(r1_desde, r1_hasta))
    r2_data = rango_seguro(rango_stats(r2_desde, r2_hasta))

    diff_total = (r1_data['total'] - r2_data['total']) if modo_rangos == 'compare' else None

    comparacion = {
        'modo': modo_rangos,
        'r1': r1_data,
        'r2': r2_data,
        'diff_total': diff_total,
        'clear_url': filter_url(r1_desde='', r1_hasta='', r2_desde='', r2_hasta=''),
    }

    if r1_desde:
        comparacion['nav_r1_ant'] = nav_url_explicit(r1_desde, r1_hasta, 1, -1)
        comparacion['nav_r1_sig'] = nav_url_explicit(r1_desde, r1_hasta, 1, 1)
    else:
        comparacion['nav_r1_ant'] = comparacion['nav_r1_sig'] = None

    if r2_desde:
        comparacion['nav_r2_ant'] = nav_url_explicit(r2_desde, r2_hasta, 2, -1)
        comparacion['nav_r2_sig'] = nav_url_explicit(r2_desde, r2_hasta, 2, 1)
    else:
        comparacion['nav_r2_ant'] = comparacion['nav_r2_sig'] = None

    # --- Lista de filtros activos (para mostrar múltiples badges) ---
    from django.contrib.auth.models import User
    filtros_list = []
    def add_filtro(key, label_val):
        remove_url = filter_url(**{key: ''})
        filtros_list.append({'key': key, 'label': label_val, 'remove_url': remove_url})

    if filtro_estado:
        add_filtro('estado', f'Estado: {ESTADO_LABELS.get(filtro_estado, filtro_estado)}')
    if filtro_tipo:
        try:
            t = Tipificacion.objects.get(slug=filtro_tipo)
            add_filtro('tipo', f'Tipo: {t.nombre}')
        except Tipificacion.DoesNotExist:
            add_filtro('tipo', f'Tipo: {filtro_tipo}')
    if filtro_direccion:
        add_filtro('direccion', f'Dirección: {filtro_direccion}')
    if filtro_tecnico_id:
        try:
            u = User.objects.get(id=filtro_tecnico_id)
            add_filtro('tecnico_id', f'Técnico: {u.get_full_name() or u.username}')
        except User.DoesNotExist:
            add_filtro('tecnico_id', f'Técnico: {filtro_tecnico_id}')
    if filtro_mes:
        m_name = meses_es[int(filtro_mes)] if filtro_mes.isdigit() and 1 <= int(filtro_mes) <= 12 else filtro_mes
        add_filtro('mes', f'Mes {m_name} {filtro_anio}')

    total_filtrado = tickets.count()

    # --- Resumen de datos para la leyenda ---
    def pct(val):
        return round((val / total_filtrado) * 100) if total_filtrado else 0

    resumen_estado = [{'label': ESTADO_LABELS.get(d['estado'], d['estado']), 'cantidad': d['cantidad'], 'porcentaje': pct(d['cantidad']), 'color': ESTADO_COLORS.get(d['estado'], '#64748b')} for d in por_estado]

    resumen_tipo = []
    for d in por_tipo:
        resumen_tipo.append({'label': d['tipo_falla'], 'cantidad': d['cantidad'], 'porcentaje': pct(d['cantidad']), 'color': TIPO_COLORS.get(d['tipo_falla'], '#64748b')})

    resumen_area = [{'label': d['direccion_area'] or '(Sin dirección)', 'cantidad': d['cantidad'], 'porcentaje': pct(d['cantidad'])} for d in por_area]

    resumen_tecnico = [{'label': d['tecnico_asignado__nombre'], 'cantidad': d['cantidad'], 'porcentaje': pct(d['cantidad'])} for d in por_tecnico]

    context = {
        'svg_estado': svg_estado,
        'svg_tipo': svg_tipo,
        'svg_area': svg_area,
        'svg_tecnico': svg_tecnico,
        'es_tecnico': is_tecnico(request.user),
        'filtros_list': filtros_list,
        'total_filtrado': total_filtrado,
        'hay_filtro': bool(filtros_list),
        'clear_url': filter_url(estado='', tipo='', direccion='', tecnico_id='', mes='', anio=''),
        'resumen_estado': resumen_estado,
        'resumen_tipo': resumen_tipo,
        'resumen_area': resumen_area,
        'resumen_tecnico': resumen_tecnico,
        'comparacion': comparacion,
    }

    return render(request, 'tickets/reportes.html', context)


@login_required
@user_passes_test(is_admin_or_tecnico)
def crear_ticket_staff(request):
    if request.method == 'POST':
        form = StaffTicketForm(request.POST)
        if form.is_valid():
            ticket = form.save(commit=False)
            if request.POST.get('atendere_yo'):
                ticket.tecnico_asignado = request.user
            ticket.save()
            form.save_m2m()
            if not ticket.tipificaciones.exists():
                analisis = analizar_ticket(ticket.descripcion)
                if analisis and analisis.get('tipos_detectados'):
                    for slug in analisis['tipos_detectados']:
                        tip, _ = Tipificacion.objects.get_or_create(
                            slug=slug,
                            defaults={'nombre': dict(Ticket.TIPO_FALLA_CHOICES).get(slug, slug)}
                        )
                        ticket.tipificaciones.add(tip)
            tiene_tipo_automatico = ticket.tipificaciones.exclude(slug='otro').exists()
            if not ticket.tecnico_asignado and tiene_tipo_automatico:
                tecnico = asignar_tecnico_round_robin()
                if tecnico:
                    ticket.tecnico_asignado = tecnico
                    ticket.save(update_fields=['tecnico_asignado'])
            if ticket.tecnico_asignado:
                _crear_notificacion(
                    f'Nuevo ticket asignado {ticket.codigo} - {ticket.solicitante}',
                    url=f'/tickets/{ticket.pk}/',
                    usuario=ticket.tecnico_asignado
                )
            if ticket.solicitante:
                _crear_notificacion_usuario(
                    f'{ticket.codigo} - Nuevo ticket creado por administración',
                    url=f'/mi-panel/ticket/{ticket.pk}/',
                    solicitante=ticket.solicitante,
                    direccion_area=ticket.direccion_area
                )
            _registrar_auditoria(request, request.user, 'CREAR', 'Ticket', id_objeto=ticket.pk, detalles=f'Staff creó el ticket {ticket.codigo} para "{ticket.solicitante}" - {ticket.descripcion[:100]}')
            messages.success(request, f'Ticket creado exitosamente. Codigo: {ticket.codigo}')
            return redirect('dashboard')
    else:
        form = StaffTicketForm()
    return render(request, 'tickets/crear_ticket_staff.html', {'form': form, 'es_admin': is_admin(request.user), 'es_tecnico': is_tecnico(request.user)})


def _build_encabezado_pdf(styles, avail_width=510):
    elements = []
    center_style = ParagraphStyle(
        'CenterTitle', parent=styles['Normal'],
        fontSize=16, leading=20, alignment=TA_CENTER,
        spaceAfter=2, textColor=colors.HexColor('#1e293b'),
        fontName='Helvetica-Bold'
    )
    sub_center_style = ParagraphStyle(
        'SubCenter', parent=styles['Normal'],
        fontSize=9, leading=12, alignment=TA_CENTER,
        spaceAfter=2, textColor=colors.HexColor('#64748b')
    )
    title = 'SOPORTE TÉCNICO'
    subtitle = 'Dirección de Tecnología'
    logo_img = None
    try:
        from django.conf import settings
        import os as _os
        logo_paths = [
            _os.path.join(settings.STATICFILES_DIRS[0], 'img', 'logo.png'),
            _os.path.join(settings.STATIC_ROOT, 'img', 'logo.png'),
        ]
        logo_full = None
        for p in logo_paths:
            if _os.path.exists(p):
                logo_full = p
                break
        if logo_full:
            logo_img = Image(logo_full, width=100, height=80)
    except Exception:
        pass

    text_col_w = avail_width - 2 * 105
    text_cell = [
        [Paragraph(title, center_style)],
        [Paragraph(subtitle, sub_center_style)],
    ]
    text_table = Table(text_cell, colWidths=[text_col_w])

    if logo_img:
        header_table = Table(
            [[logo_img, text_table]],
            colWidths=[105, text_col_w],
            hAlign='LEFT',
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 1),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
        ]))
        elements.append(header_table)
    else:
        elements.append(Paragraph(title, center_style))
        elements.append(Paragraph(subtitle, sub_center_style))
    elements.append(Spacer(1, 3*mm))
    elements.append(Table([['']], colWidths=[avail_width], rowHeights=[1],
                          style=TableStyle([('LINEBELOW', (0, 0), (-1, -1), 1, colors.HexColor('#3b82f6'))])))
    elements.append(Spacer(1, 4*mm))
    return elements


def _build_pie_pdf(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 8)
    canvas.setFillColor(colors.HexColor('#64748b'))
    pw = doc.pagesize[0]
    lm = doc.leftMargin
    rm = doc.rightMargin
    canvas.drawCentredString(pw / 2, 15, 'Dirección de Tecnología - Sistema de Tickets de Soporte Técnico')
    canvas.drawString(lm, 6, f'Generado el {timezone.now().strftime("%d/%m/%Y %H:%M")}')
    canvas.drawRightString(pw - rm, 6, f'Pág. {canvas.getPageNumber()}')
    canvas.restoreState()


@login_required
@user_passes_test(is_admin_or_tecnico)
def ticket_pdf(request, ticket_id):
    ticket = get_object_or_404(Ticket, pk=ticket_id)
    if is_tecnico(request.user) and ticket.tecnico_asignado != request.user:
        messages.error(request, 'No tiene permiso para ver este ticket.')
        return redirect('dashboard')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            topMargin=20*mm, bottomMargin=20*mm,
                            leftMargin=15*mm, rightMargin=15*mm)
    styles = getSampleStyleSheet()
    elements = []

    avail_width = A4[0] - 30*mm
    elements += _build_encabezado_pdf(styles, avail_width=avail_width)

    title_style = ParagraphStyle('TitleDoc', parent=styles['Normal'],
                                  fontSize=13, leading=16, alignment=TA_LEFT,
                                  spaceAfter=6, textColor=colors.HexColor('#1e293b'),
                                  fontName='Helvetica-Bold')
    elements.append(Paragraph(f'{ticket.codigo} - {ticket.solicitante}', title_style))
    elements.append(Spacer(1, 3*mm))

    data = [
        ['Campo', 'Valor'],
        ['Código', ticket.codigo],
        ['Solicitante', ticket.solicitante],
        ['Contacto', ticket.contacto or 'N/A'],
        ['Dirección/Departamento', ticket.direccion_area],
        ['Tipificación', ticket.nombres_tipificacion],
        ['Estado', ticket.get_estado_display()],
        ['Técnico asignado', ticket.nombre_tecnico],
        ['Fecha de creación', ticket.fecha_creacion.strftime('%d/%m/%Y %H:%M')],
        ['Última actualización', ticket.fecha_actualizacion.strftime('%d/%m/%Y %H:%M')],
    ]

    col_widths = [120, 370]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    style_table = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3b82f6')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8.5),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#f8fafc')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
    ])
    table.setStyle(style_table)
    elements.append(table)
    elements.append(Spacer(1, 4*mm))

    desc_style = ParagraphStyle('Desc', parent=styles['Normal'],
                                 fontSize=9, leading=13,
                                 spaceAfter=3, textColor=colors.HexColor('#1e293b'),
                                 fontName='Helvetica-Bold')
    val_style = ParagraphStyle('Val', parent=styles['Normal'],
                                fontSize=9, leading=13,
                                spaceAfter=6, textColor=colors.HexColor('#334155'))
    elements.append(Paragraph('Descripción:', desc_style))
    elements.append(Paragraph(ticket.descripcion, val_style))

    if ticket.notas:
        elements.append(Paragraph('Notas de seguimiento:', desc_style))
        elements.append(Paragraph(ticket.notas, val_style))

    doc.build(elements, onFirstPage=_build_pie_pdf, onLaterPages=_build_pie_pdf)
    buffer.seek(0)
    return HttpResponse(buffer, content_type='application/pdf')


@login_required
@user_passes_test(is_admin)
def tickets_pdf(request):
    tickets = Ticket.objects.all().select_related('tecnico_asignado')

    if is_tecnico(request.user):
        tickets = tickets.filter(tecnico_asignado=request.user)

    filtro_estado = request.GET.get('estado')
    filtro_tecnico = request.GET.get('tecnico')
    buscar = request.GET.get('buscar')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    if filtro_estado:
        tickets = tickets.filter(estado=filtro_estado)
    if filtro_tecnico:
        tickets = tickets.filter(tecnico_asignado_id=filtro_tecnico)
    if buscar:
        tickets = tickets.filter(
            Q(solicitante__icontains=buscar) |
            Q(descripcion__icontains=buscar) |
            Q(direccion_area__icontains=buscar)
        )
    if fecha_desde:
        tickets = tickets.filter(fecha_creacion__gte=fecha_desde)
    if fecha_hasta:
        tickets = tickets.filter(fecha_creacion__lte=fecha_hasta)

    tickets = tickets.order_by('-fecha_creacion')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                            topMargin=15*mm, bottomMargin=20*mm,
                            leftMargin=12*mm, rightMargin=12*mm)
    styles = getSampleStyleSheet()
    elements = []

    avail_width = landscape(A4)[0] - 24*mm
    elements += _build_encabezado_pdf(styles, avail_width=avail_width)

    title_style = ParagraphStyle('TitleDoc', parent=styles['Normal'],
                                  fontSize=12, leading=15, alignment=TA_LEFT,
                                  spaceAfter=6, textColor=colors.HexColor('#1e293b'),
                                  fontName='Helvetica-Bold')
    user_filter = ''
    if buscar:
        user_filter += f'  |  Búsqueda: {buscar}'
    if filtro_estado:
        user_filter += f'  |  Estado: {dict(Ticket.ESTADO_CHOICES).get(filtro_estado, filtro_estado)}'
    elements.append(Paragraph(f'Lista de Tickets ({tickets.count()} registros){user_filter}', title_style))
    elements.append(Spacer(1, 3*mm))

    header = ['Código', 'Solicitante', 'Tipificación', 'Área', 'Técnico', 'Estado', 'Fecha']
    data = [header]
    for t in tickets:
        data.append([
            t.codigo,
            t.solicitante,
            t.nombres_tipificacion,
            t.direccion_area,
            t.nombre_tecnico,
            t.get_estado_display(),
            t.fecha_creacion.strftime('%d/%m/%Y'),
        ])

    avail_width = landscape(A4)[0] - 24*mm

    header_font = 'Helvetica-Bold'
    body_font = 'Helvetica'
    header_font_size = 7.5
    body_font_size = 7
    cell_padding = 14

    col_widths = []
    for col_idx in range(len(header)):
        max_width = 0
        for row_idx, row in enumerate(data):
            font = header_font if row_idx == 0 else body_font
            font_size = header_font_size if row_idx == 0 else body_font_size
            text = str(row[col_idx])
            w = stringWidth(text, font, font_size)
            if w > max_width:
                max_width = w
        col_widths.append(max_width + cell_padding)

    total_width = sum(col_widths)
    scale = avail_width / total_width
    col_widths = [w * scale for w in col_widths]

    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign='LEFT')

    style_table = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3b82f6')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 7.5),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
    ])
    table.setStyle(style_table)
    elements.append(table)

    doc.build(elements, onFirstPage=_build_pie_pdf, onLaterPages=_build_pie_pdf)
    buffer.seek(0)
    return HttpResponse(buffer, content_type='application/pdf')


@login_required
@user_passes_test(is_admin)
def tickets_excel(request):
    tickets = Ticket.objects.all().select_related('tecnico_asignado')

    filtro_estado = request.GET.get('estado')
    filtro_tecnico = request.GET.get('tecnico')
    buscar = request.GET.get('buscar')
    fecha_desde = request.GET.get('fecha_desde')
    fecha_hasta = request.GET.get('fecha_hasta')

    if filtro_estado:
        tickets = tickets.filter(estado=filtro_estado)
    if filtro_tecnico:
        tickets = tickets.filter(tecnico_asignado_id=filtro_tecnico)
    if buscar:
        tickets = tickets.filter(
            Q(solicitante__icontains=buscar) |
            Q(descripcion__icontains=buscar) |
            Q(direccion_area__icontains=buscar)
        )
    if fecha_desde:
        tickets = tickets.filter(fecha_creacion__gte=fecha_desde)
    if fecha_hasta:
        tickets = tickets.filter(fecha_creacion__lte=fecha_hasta)

    tickets = tickets.order_by('-fecha_creacion')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tickets'

    headers = ['Código', 'Solicitante', 'Contacto', 'Dirección/Departamento',
               'Tipificación', 'Estado', 'Técnico Asignado', 'Fecha Creación',
               'Última Actualización', 'Descripción', 'Notas']

    blue_fill = PatternFill(start_color='3b82f6', end_color='3b82f6', fill_type='solid')
    white_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='e2e8f0'),
        right=Side(style='thin', color='e2e8f0'),
        top=Side(style='thin', color='e2e8f0'),
        bottom=Side(style='thin', color='e2e8f0'),
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = blue_fill
        cell.font = white_font
        cell.alignment = header_alignment
        cell.border = thin_border

    alt_fill = PatternFill(start_color='f8fafc', end_color='f8fafc', fill_type='solid')
    body_font = Font(name='Calibri', size=11)
    body_alignment = Alignment(vertical='center', wrap_text=True)

    for row_idx, t in enumerate(tickets, start=2):
        values = [
            t.codigo,
            t.solicitante,
            t.contacto or '',
            t.direccion_area,
            t.nombres_tipificacion,
            t.get_estado_display(),
            t.nombre_tecnico or '',
            t.fecha_creacion.strftime('%d/%m/%Y %H:%M'),
            t.fecha_actualizacion.strftime('%d/%m/%Y %H:%M'),
            t.descripcion or '',
            t.notas or '',
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = body_font
            cell.alignment = body_alignment
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    for col_idx in range(1, len(headers) + 1):
        max_length = 0
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 4, 60)

    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{ws.max_row}'
    ws.freeze_panes = 'A2'

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="tickets.xlsx"'
    wb.save(response)
    return response


from django.http import JsonResponse


@login_required
@user_passes_test(is_admin_or_tecnico)
def notificaciones_json(request):
    notifs = Notificacion.objects.filter(
        usuario=request.user,
        leido=False
    )[:20]
    data = [{
        'id': n.id,
        'mensaje': n.mensaje,
        'url': n.url,
        'creado_en': n.creado_en.isoformat(),
    } for n in notifs]
    return JsonResponse({'notificaciones': data, 'total': len(data)})


@login_required
@user_passes_test(is_admin_or_tecnico)
def marcar_notificacion_leida(request, notif_id):
    Notificacion.objects.filter(pk=notif_id, usuario=request.user).update(leido=True)
    return JsonResponse({'ok': True})


def notificaciones_usuario_json(request):
    nombre = request.session.get('usuario_nombre')
    area = request.session.get('usuario_area')
    if not nombre or not area:
        return JsonResponse({'notificaciones': [], 'total': 0})
    notifs = NotificacionUsuario.objects.filter(
        solicitante__iexact=nombre,
        direccion_area__iexact=area,
        leido=False
    )[:20]
    data = [{
        'id': n.id,
        'mensaje': n.mensaje,
        'url': n.url,
        'creado_en': n.creado_en.isoformat(),
    } for n in notifs]
    return JsonResponse({'notificaciones': data, 'total': len(data)})


def marcar_notificacion_usuario_leida(request, notif_id):
    nombre = request.session.get('usuario_nombre')
    area = request.session.get('usuario_area')
    if not nombre or not area:
        return JsonResponse({'ok': False})
    NotificacionUsuario.objects.filter(
        pk=notif_id,
        solicitante__iexact=nombre,
        direccion_area__iexact=area
    ).update(leido=True)
    return JsonResponse({'ok': True})


import threading


def _generar_pdf_cache():
    from django.conf import settings
    from django.template.loader import render_to_string
    from datetime import datetime
    from pathlib import Path

    cache_dir = Path(settings.BASE_DIR) / 'pdf_cache'
    cache_dir.mkdir(exist_ok=True)
    hoy = datetime.now().strftime('%Y%m%d')
    cache_file = cache_dir / f'manual_usuario_{hoy}.pdf'
    if cache_file.exists():
        return

    try:
        from weasyprint import HTML
    except ImportError:
        return

    for old in cache_dir.glob('manual_usuario_*.pdf'):
        old.unlink(missing_ok=True)

    fecha_actual = datetime.now()
    meses_es = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
                'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    fecha_legible = f"{fecha_actual.day} de {meses_es[fecha_actual.month - 1]} de {fecha_actual.year}"
    ctx = {'fecha_generacion': fecha_legible, 'anio_actual': fecha_actual.year}
    html_str = render_to_string('tickets/manual_usuario_pdf.html', ctx)
    pdf = HTML(string=html_str, base_url=str(settings.BASE_DIR / 'static')).write_pdf()
    cache_file.write_bytes(pdf)


def ayuda_manual(request):
    es_admin = request.user.is_authenticated and is_admin(request.user)
    if es_admin:
        thread = threading.Thread(target=_generar_pdf_cache, daemon=True)
        thread.start()
    return render(request, 'tickets/manual_usuario.html', {
        'es_admin': es_admin,
    })


@login_required
@user_passes_test(is_admin)
def manual_pdf(request):
    from django.conf import settings
    from datetime import datetime
    from pathlib import Path

    cache_dir = Path(settings.BASE_DIR) / 'pdf_cache'
    cache_dir.mkdir(exist_ok=True)
    hoy = datetime.now().strftime('%Y%m%d')
    cache_file = cache_dir / f'manual_usuario_{hoy}.pdf'

    if not cache_file.exists():
        _generar_pdf_cache()

    if cache_file.exists():
        pdf = cache_file.read_bytes()
        response = HttpResponse(pdf, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="Manual_de_Usuario.pdf"'
        return response

    messages.error(request, 'No se pudo generar el PDF. Intente nuevamente.')
    return redirect('ayuda_manual')


def ayuda_guia_rapida(request):
    return render(request, 'tickets/guia_rapida.html')


def ayuda_acerca_de(request):
    return render(request, 'tickets/acerca_de.html')


@login_required
@user_passes_test(is_admin)
def auditoria_list(request):
    registros = AuditLog.objects.all().select_related('usuario')
    accion_filtro = request.GET.get('accion')
    buscar = request.GET.get('buscar')
    if accion_filtro:
        registros = registros.filter(accion=accion_filtro)
    if buscar:
        registros = registros.filter(
            Q(nombre_completo__icontains=buscar) |
            Q(detalles__icontains=buscar) |
            Q(direccion_ip__icontains=buscar) |
            Q(direccion_area__icontains=buscar)
        )
    paginator = Paginator(registros, 15)
    page = request.GET.get('page')
    registros_page = paginator.get_page(page)
    context = {
        'registros': registros_page,
        'accion_filtro': accion_filtro,
        'buscar': buscar,
        'acciones': AuditLog.ACCIONES,
    }
    return render(request, 'tickets/auditoria_list.html', context)


@login_required
@user_passes_test(is_admin)
def auditoria_detalle(request, pk):
    registro = get_object_or_404(AuditLog, pk=pk)
    return render(request, 'tickets/auditoria_detalle.html', {'registro': registro})
